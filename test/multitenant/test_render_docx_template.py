"""Pin de la dépendance docxtpl + rendu réel du template Cyllene (2026-09-01).

Le tool RenderDocxTemplate (concours PDG « procédure fiche S3 ») importe
docxtpl paresseusement — la démo marchait sur le Mac (pip local) mais la
dépendance n'était PAS dans pyproject : l'image prod ne l'avait pas et le
tool serait mort au premier rendu sur le cluster. Doctrine memory
« graceful degradation = no-op silencieux » : chaque dépendance à import
paresseux exige un test statique de présence.
"""
import os
import pathlib
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

REPO = pathlib.Path(__file__).resolve().parents[2]


def test_docxtpl_declared_in_pyproject():
    assert '"docxtpl' in (REPO / "pyproject.toml").read_text(), (
        "docxtpl retiré de pyproject — le tool RenderDocxTemplate mourra en prod "
        "au premier rendu (import paresseux, aucun échec au boot)"
    )


def test_docxtpl_importable():
    from docxtpl import DocxTemplate  # noqa: F401


def test_cyllene_template_renders_end_to_end(tmp_path):
    """Rendu réel du template du concours PDG avec le schéma documenté
    (assets/templates/README.md) — attrape aussi une casse du template."""
    from docxtpl import DocxTemplate

    tpl_path = REPO / "assets" / "templates" / "cyllene_procedure.docx"
    assert tpl_path.exists(), "template cyllene_procedure.docx manquant"

    context = {
        "procedure_name": "Embauche fiche S3",
        "beneficiaires": ["DRH", "Manager hiérarchique", "Responsable sécurité"],
        "finalite": "Sécuriser le processus d'habilitation défense.",
        "revision_date": "01/09/2026",
        "revision_number": "V1",
        "revision_auteur": "Test automatique",
        "revision_approbateur": "PDG",
        "revision_nature": "Création",
        "sections": [
            {"titre": "VERIFICATIONS PREALABLES",
             "intro": "Avant de lancer la demande :",
             "puces": ["Vérifier l'identité", "Confirmer la fiche S3"]},
            {"titre": "CONSTITUTION DU DOSSIER",
             "intro": "Le dossier comprend :",
             "puces": ["Formulaire 94A complété"]},
        ],
    }
    tpl = DocxTemplate(str(tpl_path))
    tpl.render(context)
    out = tmp_path / "rendu.docx"
    tpl.save(str(out))
    assert out.stat().st_size > 10_000, "rendu suspicieusement petit"

    # Le contenu injecté doit être présent dans le document rendu
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    tables_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    rendered = text + "\n" + tables_text
    for needle in ("Embauche fiche S3", "VERIFICATIONS PREALABLES", "Formulaire 94A complété"):
        assert needle in rendered, f"« {needle} » absent du document rendu"
    assert "{{" not in rendered and "{%" not in rendered, "placeholders Jinja non rendus"
