"""
Convert the Cyllene PROCEDURE - Modèle.docx (placeholder content) into a
docxtpl-compatible Jinja template, preserving all styles, branding, fonts,
table formatting and the header/footer.
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SRC = "assets/templates/_source_cyllene_procedure.docx"
DST = sys.argv[1] if len(sys.argv) > 1 else "/tmp/cyllene_template_jinja.docx"


def _replace_runs_with_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for r in paragraph.runs[1:]:
        r.text = ""


def _insert_directive_paragraph_before(target_paragraph, directive: str):
    new_p = copy.deepcopy(target_paragraph._element)
    for r in list(new_p.findall(qn("w:r"))):
        new_p.remove(r)
    target_paragraph._element.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, target_paragraph._parent)
    p.add_run(directive)
    return p


def _insert_directive_paragraph_after(target_paragraph, directive: str):
    new_p = copy.deepcopy(target_paragraph._element)
    for r in list(new_p.findall(qn("w:r"))):
        new_p.remove(r)
    target_paragraph._element.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, target_paragraph._parent)
    p.add_run(directive)
    return p


def _delete_paragraph(paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    p._p = p._element = None


def main():
    print(f"[load] {SRC}")
    doc = Document(SRC)

    def find_para(predicate):
        for p in doc.paragraphs:
            if predicate(p):
                return p
        return None

    # 1. Title
    title = find_para(lambda p: p.text.strip().lower() == "nom de la procedure")
    if title is None:
        raise RuntimeError("Title paragraph 'Nom de la Procedure' not found.")
    _replace_runs_with_text(title, "{{ procedure_name }}")
    print("[ok] title -> {{ procedure_name }}")

    # 2. Locate all 3 'sujet n°' headings.
    section_headings = [
        p for p in doc.paragraphs
        if p.style.name == "Heading 1" and "sujet n" in p.text.lower()
    ]
    if len(section_headings) < 3:
        raise RuntimeError(f"Expected 3 'sujet n°' headings, found {len(section_headings)}.")

    # Delete sections 2 + 3 (placeholder duplicates) BEFORE wrapping section 1
    # in loops, so paragraph anchors remain stable. Compare on _element
    # because doc.paragraphs returns fresh wrappers each iteration.
    sec2_elem = section_headings[1]._element
    to_delete = []
    in_dup = False
    for p in doc.paragraphs:
        if p._element is sec2_elem:
            in_dup = True
        if in_dup:
            to_delete.append(p)
    for p in to_delete:
        _delete_paragraph(p)
    print(f"[ok] deleted {len(to_delete)} duplicate paragraphs (sections 2 & 3)")

    # 3. Section 1 -> loop body. The intro is the FIRST 'Texte Titre 1'
    # paragraph after the heading; subsequent ones are bullets. Source-text
    # matching on the colon was brittle (nbsp \xa0 in the original).
    sec1_heading = find_para(
        lambda p: p.style.name == "Heading 1" and "sujet n" in p.text.lower()
    )
    if not sec1_heading:
        raise RuntimeError("Section 1 heading not found after deletion.")
    titre1_paras = [p for p in doc.paragraphs if p.style.name == "Texte Titre 1"]
    if len(titre1_paras) < 2:
        raise RuntimeError(
            f"Expected >=2 'Texte Titre 1' paragraphs (intro + >=1 bullet), got {len(titre1_paras)}."
        )
    intro = titre1_paras[0]
    bullets = titre1_paras[1:]

    first_bullet = bullets[0]
    for b in bullets[1:]:
        _delete_paragraph(b)

    _replace_runs_with_text(sec1_heading, "{{ loop.index }} {{ s.titre }}")
    _replace_runs_with_text(intro, "{{ s.intro }}")
    _replace_runs_with_text(first_bullet, "{{ puce }}")

    _insert_directive_paragraph_before(sec1_heading, "{%p for s in sections %}")
    _insert_directive_paragraph_before(first_bullet, "{%p for puce in s.puces %}")
    _insert_directive_paragraph_after(first_bullet, "{%p endfor %}")  # close puces

    puces_endfor = first_bullet._element.getnext()
    from docx.text.paragraph import Paragraph
    puces_endfor_para = Paragraph(puces_endfor, first_bullet._parent)
    _insert_directive_paragraph_after(puces_endfor_para, "{%p endfor %}")  # close sections
    print("[ok] section loop wired (sections + puces, intro preserved)")

    # 4. Table 0: FINALITE ET BENEFICIAIRES
    t0 = doc.tables[0]
    benef_cell = t0.cell(2, 0)
    benef_paras = [p for p in benef_cell.paragraphs if p.text.strip()]
    if benef_paras:
        first_b = benef_paras[0]
        for extra in benef_paras[1:]:
            _delete_paragraph(extra)
        _replace_runs_with_text(first_b, "- {{ b }}")
        _insert_directive_paragraph_before(first_b, "{%p for b in beneficiaires %}")
        _insert_directive_paragraph_after(first_b, "{%p endfor %}")
        print("[ok] table beneficiaires -> for/endfor on bullet")

    fin_cell = t0.cell(4, 0)
    fin_paras = [p for p in fin_cell.paragraphs if p.text.strip()]
    if fin_paras:
        _replace_runs_with_text(fin_paras[0], "- {{ finalite }}")
        print("[ok] table finalite -> - {{ finalite }}")

    # 5. Table 1: SUIVI DES MODIFICATIONS
    # The "Suivi" table normally has ONE entry for a freshly-generated
    # procedure (V1 / Création). docxtpl's `{%tr for %}` pattern is brittle
    # when both `for` and `endfor` directives sit in the same row (the
    # regex matches the LAST directive greedily and erases the row's
    # content), and adding wrap-rows pollutes the visible Word template.
    # Pragmatic fallback: substitute one revision row with simple fields.
    # If a future version of the template needs an N-revision history,
    # introduce dedicated wrap-rows around row 2 instead.
    t1 = doc.tables[1]
    data_row_cells = list(t1.rows[2].cells)
    seen = set()
    unique_cells = []
    for c in data_row_cells:
        if id(c._tc) not in seen:
            unique_cells.append(c)
            seen.add(id(c._tc))
    field_map = {0: "{{ revision_date }}", 1: "{{ revision_number }}",
                 2: "{{ revision_auteur }}", 3: "{{ revision_approbateur }}",
                 4: "{{ revision_nature }}"}
    for i, cell in enumerate(unique_cells):
        if i not in field_map:
            continue
        ps = [p for p in cell.paragraphs if p.text.strip()]
        if not ps:
            continue
        _replace_runs_with_text(ps[0], field_map[i])
    print("[ok] table SUIVI DES MODIFICATIONS -> single-revision substitution (no loop)")

    Path(DST).parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)
    print(f"[save] {DST}")
    print()
    print("=== Verification: rendered text after templating ===")
    d2 = Document(DST)
    for i, p in enumerate(d2.paragraphs):
        if p.text.strip():
            print(f"  [{i}] {p.text!r}")
    print()
    for ti, t in enumerate(d2.tables):
        print(f"Table {ti}:")
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip() for c in row.cells]
            print(f"  row {ri}: {cells}")


if __name__ == "__main__":
    main()
